"""技术路线图生成模块.

基于对 923 套技术路线图模板（PPTX/DOCX/VSDX/PDF）的分析，提取出 5 种主流结构模式，
提供模式识别、模板推荐、路线图生成与多格式导出能力。

5 种结构模式（按出现频率）：
    1. 线性流程型（40%）：提出问题 → 分析问题 → 解决问题
    2. 五段式论文结构（30%）：绪论 → 理论基础 → 现状分析 → 实证研究 → 结论建议
    3. 分支并行型（15%）：研究内容 / 研究方法 / 研究思路 三线并行
    4. 层级递进型（10%）：第一级基础 → 第二级框架 → 第三级精髓
    5. 时间序列型（5%）：准备阶段 → 干预阶段 → 验收阶段

支持导出为 Draw.io XML、Mermaid 流程图、PPTX、DOCX 四种格式。

典型使用流程::

    generator = RoadmapGenerator()
    pattern = generator.recommend_pattern("journal", "经济管理", "实证研究")
    roadmap = generator.generate(outline, pattern)
    xml = generator.to_drawio_xml(roadmap)
    generator.to_pptx(roadmap, "output/roadmap.pptx")
"""

from __future__ import annotations

import logging
import re
from enum import Enum
from pathlib import Path
from typing import Any, Optional

from pydantic import BaseModel, Field

logger = logging.getLogger(__name__)

__all__ = [
    "RoadmapPattern",
    "RoadmapNode",
    "RoadmapEdge",
    "Roadmap",
    "RoadmapTemplate",
    "RoadmapGenerator",
    "ROADMAP_TEMPLATES",
    "NODE_STYLES",
]


# ========== 数据模型 ==========


class RoadmapPattern(str, Enum):
    """技术路线图结构模式枚举."""

    LINEAR = "linear"           # 线性流程型
    FIVE_STAGE = "five_stage"   # 五段式论文结构
    BRANCHING = "branching"     # 分支并行型
    HIERARCHICAL = "hierarchical"  # 层级递进型
    TEMPORAL = "temporal"       # 时间序列型


class RoadmapNode(BaseModel):
    """路线图节点.

    Attributes:
        id: 节点唯一标识.
        label: 节点显示文本.
        node_type: 节点类型 (start/process/decision/end/data).
        level: 层级，用于布局定位与树形结构.
        children: 子节点 ID 列表.
        description: 节点描述.
    """

    id: str
    label: str
    node_type: str = "process"
    level: int = 0
    children: list[str] = Field(default_factory=list)
    description: str = ""


class RoadmapEdge(BaseModel):
    """路线图边.

    Attributes:
        from_node: 起始节点 ID.
        to_node: 目标节点 ID.
        edge_type: 边类型 (sequence/branch/feedback/dashed).
        label: 边标签.
    """

    from_node: str
    to_node: str
    edge_type: str = "sequence"
    label: str = ""


class Roadmap(BaseModel):
    """技术路线图.

    Attributes:
        title: 路线图标题.
        pattern: 结构模式.
        nodes: 节点列表.
        edges: 边列表.
        description: 路线图说明.
    """

    title: str
    pattern: RoadmapPattern
    nodes: list[RoadmapNode] = Field(default_factory=list)
    edges: list[RoadmapEdge] = Field(default_factory=list)
    description: str = ""


class RoadmapTemplate(BaseModel):
    """路线图模板.

    Attributes:
        name: 模板名称.
        pattern: 结构模式.
        discipline: 适用学科领域.
        description: 模板描述.
        structure: 结构描述.
        suitability: 适用场景.
    """

    name: str
    pattern: RoadmapPattern
    discipline: str
    description: str
    structure: str
    suitability: str


# ========== 节点样式 ==========

#: 节点类型样式定义（形状 / 填充色 / 字体色）
NODE_STYLES: dict[str, dict[str, str]] = {
    "start": {"shape": "rounded", "fill": "#4CAF50", "font_color": "#FFFFFF"},
    "process": {"shape": "rectangle", "fill": "#2196F3", "font_color": "#FFFFFF"},
    "decision": {"shape": "diamond", "fill": "#FF9800", "font_color": "#FFFFFF"},
    "end": {"shape": "rounded", "fill": "#F44336", "font_color": "#FFFFFF"},
    "data": {"shape": "parallelogram", "fill": "#9C27B0", "font_color": "#FFFFFF"},
}

#: Draw.io 形状样式前缀
_DRAWIO_SHAPES: dict[str, str] = {
    "start": "rounded=1;whiteSpace=wrap;html=1;arcSize=40;",
    "process": "whiteSpace=wrap;html=1;",
    "decision": "rhombus;whiteSpace=wrap;html=1;",
    "end": "rounded=1;whiteSpace=wrap;html=1;arcSize=40;",
    "data": "shape=parallelogram;perimeter=parallelogramPerimeter;whiteSpace=wrap;html=1;",
}

#: Mermaid 节点形状模板（label 占位符）
_MERMAID_SHAPES: dict[str, str] = {
    "start": '(["{label}"])',
    "process": '["{label}"]',
    "decision": '{{"{label}"}}',
    "end": '(["{label}"])',
    "data": '[/"{label}"/]',
}

#: 布局常量
NODE_WIDTH = 160
NODE_HEIGHT = 60
COL_SPACING = 80   # 列间距
ROW_SPACING = 60   # 行间距


# ========== 模板库 ==========

ROADMAP_TEMPLATES: list[RoadmapTemplate] = [
    # ---- 线性流程型（6 个）----
    RoadmapTemplate(
        name="通用学术论文路线图",
        pattern=RoadmapPattern.LINEAR,
        discipline="通用",
        description="适用于大多数学术论文的通用线性流程路线图。",
        structure="提出问题 → 文献回顾 → 研究设计 → 数据分析 → 结论建议",
        suitability="各类期刊论文、学位论文的通用研究路线展示。",
    ),
    RoadmapTemplate(
        name="实证研究路线图",
        pattern=RoadmapPattern.LINEAR,
        discipline="经济学/管理学",
        description="面向实证研究的线性路线图，强调假设检验与数据分析流程。",
        structure="理论分析 → 假设提出 → 变量界定 → 数据收集 → 计量检验 → 稳健性检验",
        suitability="经济学、金融学、管理学等实证类研究。",
    ),
    RoadmapTemplate(
        name="案例研究路线图",
        pattern=RoadmapPattern.LINEAR,
        discipline="管理学/社会学",
        description="案例研究的线性路线图，突出案例选择与理论构建过程。",
        structure="研究问题 → 案例选择 → 数据收集 → 案内分析 → 跨案例比较 → 理论构建",
        suitability="质性研究、单案例与多案例比较研究。",
    ),
    RoadmapTemplate(
        name="理论研究路线图",
        pattern=RoadmapPattern.LINEAR,
        discipline="理论学科",
        description="理论研究的线性路线图，侧重概念推演与逻辑论证。",
        structure="概念界定 → 命题提出 → 逻辑推演 → 模型构建 → 理论验证 → 结论",
        suitability="纯理论推演、概念框架构建类研究。",
    ),
    RoadmapTemplate(
        name="文献综述路线图",
        pattern=RoadmapPattern.LINEAR,
        discipline="通用",
        description="文献综述的线性路线图，体现检索、筛选、分析与综合流程。",
        structure="主题界定 → 文献检索 → 筛选评鉴 → 分类归纳 → 综合评述 → 研究展望",
        suitability="综述类论文、文献计量分析。",
    ),
    RoadmapTemplate(
        name="实验研究路线图",
        pattern=RoadmapPattern.LINEAR,
        discipline="心理学/理工",
        description="实验研究的线性路线图，涵盖实验设计与验证全过程。",
        structure="假设提出 → 实验设计 → 被试招募 → 实验实施 → 数据处理 → 假设检验",
        suitability="心理学实验、控制实验、理工科实验研究。",
    ),
    # ---- 五段式论文结构（3 个）----
    RoadmapTemplate(
        name="经管类论文路线图",
        pattern=RoadmapPattern.FIVE_STAGE,
        discipline="经济管理",
        description="经管类学位论文经典的五段式路线图。",
        structure="绪论 → 理论基础与文献综述 → 现状分析 → 实证研究 → 结论与建议",
        suitability="经济管理类硕士/博士学位论文。",
    ),
    RoadmapTemplate(
        name="社科类论文路线图",
        pattern=RoadmapPattern.FIVE_STAGE,
        discipline="社会科学",
        description="社会科学类论文的五段式路线图，兼顾理论与社会调查。",
        structure="绪论 → 理论框架 → 现状与问题 → 调查实证 → 结论与对策",
        suitability="社会学、政治学、教育学等社科类论文。",
    ),
    RoadmapTemplate(
        name="理工科论文路线图",
        pattern=RoadmapPattern.FIVE_STAGE,
        discipline="理工",
        description="理工科论文的五段式路线图，突出方法与实验验证。",
        structure="绪论 → 理论方法 → 系统设计 → 实验验证 → 结论与展望",
        suitability="工学、理学等理工科研究论文。",
    ),
    # ---- 分支并行型（2 个）----
    RoadmapTemplate(
        name="多方法并行研究路线图",
        pattern=RoadmapPattern.BRANCHING,
        discipline="通用",
        description="多种研究方法并行的分支路线图。",
        structure="研究问题 → 研究内容 / 研究方法 / 研究思路 三线并行 → 综合结论",
        suitability="混合方法研究、多维度并行研究。",
    ),
    RoadmapTemplate(
        name="跨学科研究路线图",
        pattern=RoadmapPattern.BRANCHING,
        discipline="跨学科",
        description="跨学科研究的分支并行路线图。",
        structure="共同问题 → 学科A视角 / 学科B视角 / 学科C视角 → 交叉融合 → 结论",
        suitability="跨学科交叉研究、多视角综合研究。",
    ),
    # ---- 层级递进型（2 个）----
    RoadmapTemplate(
        name="理论构建路线图",
        pattern=RoadmapPattern.HIERARCHICAL,
        discipline="理论研究",
        description="理论构建的层级递进路线图，由基础到精髓逐层深入。",
        structure="第一级：基础概念 → 第二级：理论框架 → 第三级：核心命题",
        suitability="理论体系构建、概念框架设计。",
    ),
    RoadmapTemplate(
        name="模型构建路线图",
        pattern=RoadmapPattern.HIERARCHICAL,
        discipline="建模研究",
        description="模型构建的层级递进路线图，由变量到模型逐层抽象。",
        structure="第一级：变量界定 → 第二级：关系框架 → 第三级：核心模型",
        suitability="计量模型构建、系统建模研究。",
    ),
    # ---- 时间序列型（2 个）----
    RoadmapTemplate(
        name="纵向研究路线图",
        pattern=RoadmapPattern.TEMPORAL,
        discipline="社会科学/医学",
        description="纵向研究的时序路线图，按时间阶段展开。",
        structure="准备阶段（基线测量）→ 干预阶段（跟踪观察）→ 验收阶段（终期评估）",
        suitability="纵向追踪研究、队列研究、干预实验。",
    ),
    RoadmapTemplate(
        name="项目周期路线图",
        pattern=RoadmapPattern.TEMPORAL,
        discipline="工程管理",
        description="项目周期的时序路线图，按阶段推进。",
        structure="准备阶段（立项规划）→ 干预阶段（实施执行）→ 验收阶段（结题评估）",
        suitability="工程项目、课题立项、研发项目管理。",
    ),
]


# ========== 核心类 ==========


class RoadmapGenerator:
    """技术路线图生成器.

    提供模式识别、模板推荐、路线图生成与多格式导出能力。基于 923 套模板分析结果，
    内置 5 种结构模式与 15 个预定义模板。

    功能：
        - 根据论文类型/学科/方法自动推荐路线图模式
        - 基于论文大纲生成结构化路线图
        - 导出 Draw.io XML / Mermaid / PPTX / DOCX

    使用示例::

        generator = RoadmapGenerator()
        pattern = generator.recommend_pattern("journal", "经济管理", "实证研究")
        roadmap = generator.generate(outline, pattern)
        generator.to_pptx(roadmap, "roadmap.pptx")
    """

    def __init__(self) -> None:
        """初始化路线图生成器."""
        self._node_counter = 0
        self._temporal_stages = 3
        logger.debug("RoadmapGenerator 初始化完成")

    # ---- 模式识别与推荐 ----

    def recommend_pattern(
        self,
        paper_type: str,
        discipline: str,
        research_method: str,
    ) -> RoadmapPattern:
        """根据论文类型、学科与研究方法自动推荐路线图模式.

        推荐优先级：研究方法 > 学科领域 > 论文类型。

        Args:
            paper_type: 论文类型 (journal/conference/thesis/review 等).
            discipline: 学科领域.
            research_method: 研究方法.

        Returns:
            推荐的 RoadmapPattern 枚举值。
        """
        rm = (research_method or "").lower()
        disc = (discipline or "").lower()
        pt = (paper_type or "").lower()

        # 时间序列型：纵向/面板/追踪/队列
        temporal_keywords = [
            "纵向", "面板", "时间序列", "追踪", "队列", "干预",
            "longitudinal", "panel", "time series", "cohort",
        ]
        if any(k in rm or k in disc for k in temporal_keywords):
            return RoadmapPattern.TEMPORAL

        # 分支并行型：多方法/混合/跨学科/并行
        branch_keywords = [
            "多方法", "混合方法", "跨学科", "并行", "多视角",
            "multi-method", "mixed", "interdisciplinary",
        ]
        if any(k in rm or k in disc for k in branch_keywords):
            return RoadmapPattern.BRANCHING

        # 层级递进型：理论构建/模型构建/框架构建
        hier_keywords = [
            "理论构建", "模型构建", "框架构建", "层级", "体系构建",
            "theory build", "model build",
        ]
        if any(k in rm for k in hier_keywords):
            return RoadmapPattern.HIERARCHICAL

        # 五段式：经管/社科/理工类学位论文
        five_stage_disciplines = [
            "经济", "管理", "经管", "社会", "社科", "理工", "工学", "理学",
            "economics", "management", "social",
        ]
        if any(k in disc for k in five_stage_disciplines):
            return RoadmapPattern.FIVE_STAGE

        # 综述类 -> 线性
        if "review" in pt or "综述" in pt or "综述" in rm:
            return RoadmapPattern.LINEAR

        # 默认线性流程型
        return RoadmapPattern.LINEAR

    def get_templates(
        self, pattern: Optional[RoadmapPattern] = None
    ) -> list[RoadmapTemplate]:
        """获取模板列表，可按结构模式筛选.

        Args:
            pattern: 结构模式，None 时返回全部模板。

        Returns:
            模板列表。
        """
        if pattern is None:
            return list(ROADMAP_TEMPLATES)
        return [t for t in ROADMAP_TEMPLATES if t.pattern == pattern]

    # ---- 大纲解析 ----

    def parse_outline(self, outline: dict) -> list[RoadmapNode]:
        """从论文大纲中提取节点.

        支持多级大纲（章 → 节 → 小节），自动判断节点类型。
        大纲格式示例::

            {
                "title": "论文标题",
                "sections": [
                    {"title": "绪论", "sections": [
                        {"title": "研究背景"},
                        {"title": "研究问题"},
                    ]},
                    {"title": "结论"},
                ]
            }

        也支持 "chapters" 键。sections 元素可为字符串或字典。

        Args:
            outline: 论文大纲字典。

        Returns:
            路线图节点列表（按文档顺序，含层级与父子关系）。
        """
        self._node_counter = 0
        nodes: list[RoadmapNode] = []
        node_map: dict[str, RoadmapNode] = {}

        sections = outline.get("sections", outline.get("chapters", []))
        self._walk_sections(sections, nodes, node_map, parent_id=None, depth=0)

        # 根据位置与关键词细化节点类型
        total = len(nodes)
        for i, node in enumerate(nodes):
            node.node_type = self._detect_node_type(node.label, i, total)
        return nodes

    def _walk_sections(
        self,
        sections: list[Any],
        nodes: list[RoadmapNode],
        node_map: dict[str, RoadmapNode],
        parent_id: Optional[str],
        depth: int,
    ) -> None:
        """递归遍历大纲章节，构建节点与父子关系."""
        for sec in sections:
            if isinstance(sec, str):
                sec = {"title": sec}
            label = sec.get("title", sec.get("name", "")) if isinstance(sec, dict) else str(sec)
            if not label:
                continue
            node_id = f"n{self._node_counter}"
            self._node_counter += 1
            node = RoadmapNode(
                id=node_id,
                label=label,
                node_type="process",
                level=depth,
                children=[],
                description=sec.get("description", "") if isinstance(sec, dict) else "",
            )
            nodes.append(node)
            node_map[node_id] = node
            if parent_id is not None and parent_id in node_map:
                node_map[parent_id].children.append(node_id)
            child_sections = (
                sec.get("sections", sec.get("children", [])) if isinstance(sec, dict) else []
            )
            if child_sections:
                self._walk_sections(child_sections, nodes, node_map, node_id, depth + 1)

    @staticmethod
    def _detect_node_type(label: str, index: int, total: int) -> str:
        """根据标签关键词与位置自动判断节点类型.

        Args:
            label: 节点文本.
            index: 节点在列表中的位置.
            total: 节点总数.

        Returns:
            节点类型 (start/process/decision/end/data)。
        """
        if index == 0:
            return "start"
        if index == total - 1:
            return "end"
        # 判断节点
        if any(k in label for k in ["是否", "选择", "判断", "?", "？", "是否需要"]):
            return "decision"
        # 数据节点
        if any(k in label for k in ["数据", "资料", "样本", "文献", "问卷", "数据收集", "data"]):
            return "data"
        return "process"

    # ---- 路线图生成 ----

    def generate(
        self,
        outline: dict,
        pattern: Optional[RoadmapPattern] = None,
    ) -> Roadmap:
        """基于论文大纲生成路线图.

        自动从大纲提取节点与边，若未指定模式则自动推荐。

        Args:
            outline: 论文大纲字典.
            pattern: 指定结构模式，None 时自动推荐。

        Returns:
            生成的 Roadmap 对象。
        """
        title = outline.get("title", "研究路线图")
        # 记录时序阶段数（如有）
        self._temporal_stages = int(outline.get("stages", 3))

        nodes = self.parse_outline(outline)

        if pattern is None:
            pattern = self.recommend_pattern(
                paper_type=outline.get("paper_type", ""),
                discipline=outline.get("discipline", ""),
                research_method=outline.get("research_method", ""),
            )

        edges = self._generate_edges(nodes, pattern)
        description = self._build_description(pattern)

        roadmap = Roadmap(
            title=title,
            pattern=pattern,
            nodes=nodes,
            edges=edges,
            description=description,
        )
        logger.info(
            "已生成路线图: %s（模式=%s，节点=%d，边=%d）",
            title, pattern.value, len(nodes), len(edges),
        )
        return roadmap

    def generate_from_spec(self, spec: dict) -> Roadmap:
        """基于论文规格文档生成路线图.

        从 PaperSpec 格式的规格字典中提取大纲与元信息，推荐模式并生成路线图。
        支持大纲为字符串（Markdown/编号文本）或字典。

        Args:
            spec: 论文规格字典，可含 topic/paper_type/discipline/outline/research_method 等字段.

        Returns:
            生成的 Roadmap 对象。
        """
        outline = spec.get("outline")
        if isinstance(outline, str):
            outline = self._parse_outline_text(outline)
        if not isinstance(outline, dict):
            outline = {"title": "研究路线图", "sections": []}
        # 标题优先取 spec.topic，其次取大纲自带标题
        topic = spec.get("topic")
        if topic:
            outline["title"] = topic
        elif "title" not in outline:
            outline["title"] = "研究路线图"

        metadata = spec.get("metadata", {}) if isinstance(spec.get("metadata"), dict) else {}
        pattern = self.recommend_pattern(
            paper_type=spec.get("paper_type", ""),
            discipline=spec.get("discipline", metadata.get("discipline", "")),
            research_method=spec.get("research_method", metadata.get("research_method", "")),
        )
        return self.generate(outline, pattern)

    def _parse_outline_text(self, text: str) -> dict:
        """将文本大纲（Markdown 标题或编号列表）解析为结构化字典.

        支持 ``# / ## / ###`` 标题与 ``1. / 1.1 / 1.1.1`` 编号两种格式。

        Args:
            text: 文本大纲.

        Returns:
            结构化大纲字典 ``{"title": ..., "sections": [...]}``.
        """
        lines = [ln.rstrip() for ln in text.splitlines() if ln.strip()]
        if not lines:
            return {"title": "研究路线图", "sections": []}

        title = "研究路线图"
        root_sections: list[dict] = []
        stack: list[tuple[int, dict]] = []

        for line in lines:
            level, label = self._parse_heading_level(line)
            section = {"title": label, "level": level, "sections": []}
            while stack and stack[-1][0] >= level:
                stack.pop()
            if not stack:
                root_sections.append(section)
            else:
                stack[-1][1]["sections"].append(section)
            stack.append((level, section))

        # 首个带子节点的顶层标题作为文档标题，其子节点提升为顶层 sections
        if (
            root_sections
            and root_sections[0].get("level", 0) == 0
            and root_sections[0].get("sections")
        ):
            title = root_sections[0]["title"]
            sections = root_sections[0]["sections"] + root_sections[1:]
        else:
            title = root_sections[0]["title"] if root_sections else "研究路线图"
            sections = root_sections
        return {"title": title, "sections": sections}

    @staticmethod
    def _parse_heading_level(line: str) -> tuple[int, str]:
        """解析单行标题的层级与文本.

        Args:
            line: 单行标题文本.

        Returns:
            (level, label) 元组，level 从 0 开始。
        """
        # Markdown 标题
        if line.lstrip().startswith("#"):
            stripped = line.lstrip()
            hashes = len(stripped) - len(stripped.lstrip("#"))
            return max(hashes - 1, 0), stripped.lstrip("#").strip()
        # 编号列表：1. / 1.1 / 1.1.1
        m = re.match(r"^(\d+(?:\.\d+)*)[.、\)]?\s*(.+)", line)
        if m:
            nums = m.group(1).split(".")
            return len(nums) - 1, m.group(2).strip()
        # 项目符号
        if line.lstrip().startswith(("-", "*", "•")):
            return 0, line.lstrip("-*• ").strip()
        return 0, line.strip()

    def _build_description(self, pattern: RoadmapPattern) -> str:
        """根据模式构建路线图说明文字."""
        descs = {
            RoadmapPattern.LINEAR: "线性流程型路线图：按研究逻辑顺序串联各环节。",
            RoadmapPattern.FIVE_STAGE: "五段式论文结构路线图：绪论 → 理论基础 → 现状分析 → 实证研究 → 结论建议。",
            RoadmapPattern.BRANCHING: "分支并行型路线图：多条研究主线并行推进后综合。",
            RoadmapPattern.HIERARCHICAL: "层级递进型路线图：由基础到精髓逐层深入。",
            RoadmapPattern.TEMPORAL: "时间序列型路线图：按准备 → 干预 → 验收阶段时序推进。",
        }
        return descs.get(pattern, "技术路线图。")

    # ---- 边与布局生成 ----

    def _generate_edges(
        self, nodes: list[RoadmapNode], pattern: RoadmapPattern
    ) -> list[RoadmapEdge]:
        """根据模式生成边."""
        edges: list[RoadmapEdge] = []
        if not nodes:
            return edges

        if pattern in (RoadmapPattern.BRANCHING, RoadmapPattern.HIERARCHICAL):
            # 基于父子关系构建树形边
            has_children = any(n.children for n in nodes)
            if has_children:
                for node in nodes:
                    for child_id in node.children:
                        etype = (
                            "branch"
                            if pattern == RoadmapPattern.BRANCHING and node.level == 0
                            else "sequence"
                        )
                        edges.append(RoadmapEdge(
                            from_node=node.id, to_node=child_id, edge_type=etype
                        ))
                return edges
            # 无父子关系时退化为顺序连接
            for i in range(1, len(nodes)):
                edges.append(RoadmapEdge(
                    from_node=nodes[i - 1].id, to_node=nodes[i].id, edge_type="sequence"
                ))
            return edges

        # 顺序型模式：分组后组内顺序、组间顺序
        groups = self._get_groups(nodes, pattern)
        prev_last: Optional[RoadmapNode] = None
        for group in groups:
            for i in range(1, len(group)):
                edges.append(RoadmapEdge(
                    from_node=group[i - 1].id, to_node=group[i].id, edge_type="sequence"
                ))
            if prev_last is not None and group:
                edges.append(RoadmapEdge(
                    from_node=prev_last.id, to_node=group[0].id, edge_type="sequence"
                ))
            if group:
                prev_last = group[-1]
        return edges

    def _get_groups(
        self, nodes: list[RoadmapNode], pattern: RoadmapPattern
    ) -> list[list[RoadmapNode]]:
        """按模式将节点分组，用于布局与边生成."""
        if not nodes:
            return []
        if pattern == RoadmapPattern.LINEAR:
            return [list(nodes)]
        if pattern == RoadmapPattern.FIVE_STAGE:
            return self._group_nodes(nodes, 5)
        if pattern == RoadmapPattern.TEMPORAL:
            return self._group_nodes(nodes, self._temporal_stages)
        # 树形：按 level 分组
        by_level: dict[int, list[RoadmapNode]] = {}
        for node in nodes:
            by_level.setdefault(node.level, []).append(node)
        return [by_level[k] for k in sorted(by_level)]

    @staticmethod
    def _group_nodes(nodes: list[RoadmapNode], n: int) -> list[list[RoadmapNode]]:
        """将节点列表连续地均分为 n 组（不足时按实际数量分组）."""
        total = len(nodes)
        if total == 0:
            return []
        n = max(1, min(n, total))
        groups: list[list[RoadmapNode]] = [[] for _ in range(n)]
        for i, node in enumerate(nodes):
            idx = min(i * n // total, n - 1)
            groups[idx].append(node)
        return groups

    def _layout_nodes(
        self, nodes: list[RoadmapNode], pattern: RoadmapPattern
    ) -> dict[str, tuple[float, float]]:
        """计算各节点的布局坐标（左上角，布局单位）."""
        positions: dict[str, tuple[float, float]] = {}
        if not nodes:
            return positions
        col_step = NODE_WIDTH + COL_SPACING
        row_step = NODE_HEIGHT + ROW_SPACING

        if pattern == RoadmapPattern.LINEAR:
            for i, node in enumerate(nodes):
                positions[node.id] = (0.0, i * row_step)
        elif pattern in (RoadmapPattern.FIVE_STAGE, RoadmapPattern.TEMPORAL):
            groups = self._get_groups(nodes, pattern)
            for col, group in enumerate(groups):
                for row, node in enumerate(group):
                    positions[node.id] = (col * col_step, row * row_step)
        else:  # BRANCHING / HIERARCHICAL 树形布局
            groups = self._get_groups(nodes, pattern)
            for row, group in enumerate(groups):
                n = len(group)
                for col, node in enumerate(group):
                    x = (col - (n - 1) / 2) * col_step
                    positions[node.id] = (x, row * row_step)
        return positions

    # ---- Draw.io XML 导出 ----

    def to_drawio_xml(self, roadmap: Roadmap) -> str:
        """生成 Draw.io 兼容的 XML 格式.

        包含节点形状、填充色、字体色与连接边，支持中文标签。

        Args:
            roadmap: 路线图对象.

        Returns:
            Draw.io mxGraphModel XML 字符串。
        """
        positions = self._layout_nodes(roadmap.nodes, roadmap.pattern)
        lines: list[str] = [
            '<mxGraphModel dx="1422" dy="757" grid="1" gridSize="10" guides="1" '
            'tooltips="1" connect="1" arrows="1" fold="1" page="1" pageScale="1" '
            'pageWidth="1169" pageHeight="827" math="0" shadow="0">',
            "  <root>",
            '    <mxCell id="0"/>',
            '    <mxCell id="1" parent="0"/>',
        ]

        # 节点
        for node in roadmap.nodes:
            style = self._drawio_node_style(node.node_type)
            value = self._xml_escape(node.label)
            x, y = positions.get(node.id, (0.0, 0.0))
            lines.append(
                f'    <mxCell id="{node.id}" value="{value}" '
                f'style="{style}" vertex="1" parent="1">'
            )
            lines.append(
                f'      <mxGeometry x="{x:.0f}" y="{y:.0f}" '
                f'width="{NODE_WIDTH}" height="{NODE_HEIGHT}" as="geometry"/>'
            )
            lines.append("    </mxCell>")

        # 边
        for i, edge in enumerate(roadmap.edges):
            style = self._drawio_edge_style(edge, positions)
            label = self._xml_escape(edge.label)
            lines.append(
                f'    <mxCell id="e{i}" value="{label}" '
                f'style="{style}" edge="1" parent="1" '
                f'source="{edge.from_node}" target="{edge.to_node}">'
            )
            lines.append('      <mxGeometry relative="1" as="geometry"/>')
            lines.append("    </mxCell>")

        lines.extend(["  </root>", "</mxGraphModel>"])
        return "\n".join(lines)

    def _drawio_node_style(self, node_type: str) -> str:
        """构建 Draw.io 节点样式字符串."""
        base = _DRAWIO_SHAPES.get(node_type, _DRAWIO_SHAPES["process"])
        style = NODE_STYLES.get(node_type, NODE_STYLES["process"])
        return (
            f"{base}fillColor={style['fill']};"
            f"fontColor={style['font_color']};strokeColor=#FFFFFF;"
        )

    def _drawio_edge_style(
        self,
        edge: RoadmapEdge,
        positions: dict[str, tuple[float, float]],
    ) -> str:
        """构建 Draw.io 边样式字符串（含出入连接点）."""
        dashed = "dashed=1;" if edge.edge_type in ("feedback", "dashed") else ""
        src = positions.get(edge.from_node)
        tgt = positions.get(edge.to_node)
        exit_entry = ""
        if src is not None and tgt is not None:
            ex, ey, enx, eny = self._edge_exit_entry(src, tgt)
            exit_entry = (
                f"exitX={ex};exitY={ey};exitDx=0;exitDy=0;"
                f"entryX={enx};entryY={eny};entryDx=0;entryDy=0;"
            )
        return (
            f"edgeStyle=orthogonalEdgeStyle;rounded=1;orthogonalLoop=1;"
            f"jettySize=auto;html=1;{dashed}{exit_entry}"
        )

    @staticmethod
    def _edge_exit_entry(
        src: tuple[float, float], tgt: tuple[float, float]
    ) -> tuple[float, float, float, float]:
        """根据源/目标相对位置计算边的出入连接点（0~1）."""
        scx, scy = src[0] + NODE_WIDTH / 2, src[1] + NODE_HEIGHT / 2
        tcx, tcy = tgt[0] + NODE_WIDTH / 2, tgt[1] + NODE_HEIGHT / 2
        dx, dy = tcx - scx, tcy - scy
        if abs(dy) >= abs(dx):
            return (0.5, 1.0, 0.5, 0.0) if dy >= 0 else (0.5, 0.0, 0.5, 1.0)
        return (1.0, 0.5, 0.0, 0.5) if dx >= 0 else (0.0, 0.5, 1.0, 0.5)

    @staticmethod
    def _xml_escape(text: str) -> str:
        """转义 XML 特殊字符."""
        return (
            text.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;")
        )

    # ---- Mermaid 导出 ----

    def to_mermaid(self, roadmap: Roadmap) -> str:
        """生成 Mermaid 流程图语法，便于在 Markdown 中展示.

        Args:
            roadmap: 路线图对象.

        Returns:
            Mermaid flowchart 语法字符串。
        """
        direction = "LR" if roadmap.pattern in (
            RoadmapPattern.FIVE_STAGE, RoadmapPattern.TEMPORAL
        ) else "TD"
        lines: list[str] = [f"flowchart {direction}"]

        # 节点定义
        for node in roadmap.nodes:
            shape = _MERMAID_SHAPES.get(node.node_type, _MERMAID_SHAPES["process"])
            label = node.label.replace('"', "'")
            lines.append(f"    {node.id}{shape.format(label=label)}")

        # 边定义
        for edge in roadmap.edges:
            arrow = self._mermaid_arrow(edge.edge_type)
            if edge.label:
                label = edge.label.replace('"', "'")
                lines.append(f"    {edge.from_node} {arrow}|{label}| {edge.to_node}")
            else:
                lines.append(f"    {edge.from_node} {arrow} {edge.to_node}")

        return "\n".join(lines)

    @staticmethod
    def _mermaid_arrow(edge_type: str) -> str:
        """根据边类型返回 Mermaid 箭头语法."""
        if edge_type in ("feedback", "dashed"):
            return "-.->"
        return "-->"

    # ---- PPTX 导出 ----

    def to_pptx(self, roadmap: Roadmap, output_path: str | Path) -> str:
        """使用 python-pptx 生成 PPTX 文件.

        节点渲染为带样式的形状，边渲染为带箭头的连接线。

        Args:
            roadmap: 路线图对象.
            output_path: 输出文件路径.

        Returns:
            生成的 PPTX 文件路径（字符串）。

        Raises:
            ImportError: python-pptx 未安装。
        """
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor
            from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
            from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        except ImportError as e:
            raise ImportError(
                "python-pptx 是 PPTX 导出所必需的依赖。请安装: pip install python-pptx"
            ) from e

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白布局

        # 标题
        title_box = slide.shapes.add_textbox(
            Inches(0.3), Inches(0.2), Inches(12.7), Inches(0.6)
        )
        tf = title_box.text_frame
        tf.text = roadmap.title
        title_box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        title_run = tf.paragraphs[0].runs[0]
        title_run.font.size = Pt(24)
        title_run.font.bold = True

        positions = self._layout_nodes(roadmap.nodes, roadmap.pattern)
        scale, offset_x, offset_y = self._compute_scale(
            positions, slide_w_in=12.5, slide_h_in=5.8
        )

        pptx_shapes: dict[str, Any] = {}
        shape_map = {
            "start": MSO_SHAPE.ROUNDED_RECTANGLE,
            "process": MSO_SHAPE.RECTANGLE,
            "decision": MSO_SHAPE.DIAMOND,
            "end": MSO_SHAPE.ROUNDED_RECTANGLE,
            "data": MSO_SHAPE.PARALLELOGRAM,
        }

        for node in roadmap.nodes:
            x, y = positions.get(node.id, (0.0, 0.0))
            left = Inches((x - offset_x) * scale + 0.4)
            top = Inches((y - offset_y) * scale + 1.0)
            width = Inches(NODE_WIDTH * scale)
            height = Inches(NODE_HEIGHT * scale)
            shape = slide.shapes.add_shape(
                shape_map.get(node.node_type, MSO_SHAPE.RECTANGLE),
                left, top, width, height,
            )
            style = NODE_STYLES.get(node.node_type, NODE_STYLES["process"])
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(style["fill"].lstrip("#"))
            shape.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shape.line.width = Pt(1)
            stf = shape.text_frame
            stf.text = node.label
            stf.word_wrap = True
            stf.paragraphs[0].alignment = PP_ALIGN.CENTER
            for para in stf.paragraphs:
                for run in para.runs:
                    run.font.color.rgb = RGBColor.from_string(style["font_color"].lstrip("#"))
                    run.font.size = Pt(11)
            stf.vertical_anchor = MSO_ANCHOR.MIDDLE
            pptx_shapes[node.id] = shape

        # 连接线
        for edge in roadmap.edges:
            src = positions.get(edge.from_node)
            tgt = positions.get(edge.to_node)
            if src is None or tgt is None:
                continue
            x1, y1, x2, y2 = self._connector_points(src, tgt)
            conn = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches((x1 - offset_x) * scale + 0.4),
                Inches((y1 - offset_y) * scale + 1.0),
                Inches((x2 - offset_x) * scale + 0.4),
                Inches((y2 - offset_y) * scale + 1.0),
            )
            conn.line.color.rgb = RGBColor(0x60, 0x60, 0x60)
            conn.line.width = Pt(1.5)
            if edge.edge_type in ("feedback", "dashed"):
                _set_line_dashed(conn)
            _add_arrowhead(conn)

        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out))
        logger.info("PPTX 已生成: %s", out)
        return str(out)

    @staticmethod
    def _connector_points(
        src: tuple[float, float], tgt: tuple[float, float]
    ) -> tuple[float, float, float, float]:
        """计算连接线起止点（节点边缘，布局单位）."""
        ex, ey, enx, eny = RoadmapGenerator._edge_exit_entry(src, tgt)
        x1 = src[0] + ex * NODE_WIDTH
        y1 = src[1] + ey * NODE_HEIGHT
        x2 = tgt[0] + enx * NODE_WIDTH
        y2 = tgt[1] + eny * NODE_HEIGHT
        return x1, y1, x2, y2

    @staticmethod
    def _compute_scale(
        positions: dict[str, tuple[float, float]],
        slide_w_in: float,
        slide_h_in: float,
    ) -> tuple[float, float, float]:
        """计算布局到幻灯片的缩放比例与偏移量.

        Args:
            positions: 节点布局坐标.
            slide_w_in: 可用宽度（英寸）.
            slide_h_in: 可用高度（英寸）.

        Returns:
            (scale, offset_x, offset_y) 元组。offset 为布局单位下内容左上角坐标。
        """
        if not positions:
            return 1.0, 0.0, 0.0
        xs = [p[0] for p in positions.values()]
        ys = [p[1] for p in positions.values()]
        min_x, max_x = min(xs), max(xs) + NODE_WIDTH
        min_y, max_y = min(ys), max(ys) + NODE_HEIGHT
        content_w = max(max_x - min_x, 1.0)
        content_h = max(max_y - min_y, 1.0)
        scale = min(slide_w_in / content_w, slide_h_in / content_h, 1.0)
        return scale, min_x, min_y

    # ---- DOCX 导出 ----

    def to_docx(self, roadmap: Roadmap, output_path: str | Path) -> str:
        """使用 python-docx 生成 DOCX 文件.

        节点渲染为表格单元格，边以箭头符号表示。

        Args:
            roadmap: 路线图对象.
            output_path: 输出文件路径.

        Returns:
            生成的 DOCX 文件路径（字符串）。

        Raises:
            ImportError: python-docx 未安装。
        """
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.oxml.ns import qn
        except ImportError as e:
            raise ImportError(
                "python-docx 是 DOCX 导出所必需的依赖。请安装: pip install python-docx"
            ) from e

        doc = Document()

        # 默认中文字体
        style = doc.styles["Normal"]
        style.font.name = "宋体"
        style.font.size = Pt(12)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

        # 标题
        doc.add_heading(roadmap.title, level=0)
        doc.add_paragraph(roadmap.description)

        # 模式信息
        info_para = doc.add_paragraph()
        info_para.add_run(f"结构模式：{roadmap.pattern.value}　|　节点数："
                          f"{len(roadmap.nodes)}　|　连接数：{len(roadmap.edges)}")
        info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 分组渲染流程表
        groups = self._get_groups(roadmap.nodes, roadmap.pattern)
        if not groups:
            doc.add_paragraph("（无节点）")
        else:
            max_nodes = max(len(g) for g in groups)
            ncols = max(max_nodes * 2 - 1, 1)
            table = doc.add_table(rows=len(groups), cols=ncols)
            table.style = "Table Grid"
            table.alignment = 1  # 居中

            for ri, group in enumerate(groups):
                for ci, node in enumerate(group):
                    cell = table.rows[ri].cells[ci * 2]
                    cell.text = node.label
                    for para in cell.paragraphs:
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in para.runs:
                            run.bold = True
                            run.font.size = Pt(10)
                    # 箭头单元格
                    if ci * 2 + 1 < ncols:
                        arrow_cell = table.rows[ri].cells[ci * 2 + 1]
                        arrow_cell.text = "→"
                        for para in arrow_cell.paragraphs:
                            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            for run in para.runs:
                                run.font.size = Pt(14)
                                run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
                # 清空多余单元格
                for ci in range(len(group) * 2 - 1, ncols):
                    table.rows[ri].cells[ci].text = ""

            # 组间流向说明
            if len(groups) > 1:
                flow_para = doc.add_paragraph()
                flow_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                flow_run = flow_para.add_run("↓ " * (len(groups) - 1))
                flow_run.font.size = Pt(14)
                flow_run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

        # 节点明细表
        doc.add_heading("节点明细", level=1)
        detail = doc.add_table(rows=1, cols=4)
        detail.style = "Table Grid"
        hdr = detail.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "序号", "节点", "类型", "说明"
        for cell in hdr:
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.bold = True
        for idx, node in enumerate(roadmap.nodes, 1):
            row = detail.add_row().cells
            row[0].text = str(idx)
            row[1].text = node.label
            row[2].text = node.node_type
            row[3].text = node.description
            for cell in row:
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))
        logger.info("DOCX 已生成: %s", out)
        return str(out)


# ========== PPTX 连接线辅助函数 ==========


def _add_arrowhead(connector: Any) -> None:
    """为连接线添加箭头头部（通过 XML 操作）."""
    try:
        from pptx.oxml.ns import qn
    except ImportError:
        return
    try:
        spPr = connector._element.spPr
        ln = spPr.find(qn("a:ln"))
        if ln is None:
            ln = spPr.makeelement(qn("a:ln"), {"w": "12700"})
            spPr.append(ln)
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    except Exception:
        logger.debug("添加箭头头部失败，忽略", exc_info=True)


def _set_line_dashed(connector: Any) -> None:
    """将连接线设为虚线."""
    try:
        from pptx.oxml.ns import qn
    except ImportError:
        return
    try:
        spPr = connector._element.spPr
        ln = spPr.find(qn("a:ln"))
        if ln is None:
            ln = spPr.makeelement(qn("a:ln"), {"w": "12700"})
            spPr.append(ln)
        dash = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
        ln.append(dash)
    except Exception:
        logger.debug("设置虚线失败，忽略", exc_info=True)
